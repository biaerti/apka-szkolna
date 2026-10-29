from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\VS-Projekty\apka-szkolna")
OUT = ROOT / "output" / "documents" / "arkusze-wyrownawcze"
OUT.mkdir(parents=True, exist_ok=True)

FONT = "Aptos"
INK = "1F2937"
BLUE = "1F4E78"
PALE_BLUE = "EAF3F8"
PALE_GRAY = "F3F4F6"
MID_GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=MID_GRAY, size="7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(p, before=0, after=4, line=1.05, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_text(p, text, size=10.5, bold=False, color=INK, italic=False):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.35)
    sec.bottom_margin = Cm(1.25)
    sec.left_margin = Cm(1.55)
    sec.right_margin = Cm(1.55)
    sec.header_distance = Cm(0.55)
    sec.footer_distance = Cm(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(3)
    title_ppr = title.element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)

    for name, size in (("Heading 1", 13), ("Heading 2", 11.5)):
        st = doc.styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(7)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def add_footer(doc, label):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(p, after=0)
    add_text(p, f"{label}  |  język polski  |  klasa IV", size=8, color="6B7280")


def add_identity(doc, number, title, focus):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, f"Arkusz {number}  {title}", size=19, bold=True, color="000000")

    p = doc.add_paragraph()
    set_paragraph(p, after=7)
    add_text(p, focus, size=10.5, color="374151")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [10.0, 4.0, 3.4]
    labels = ["Imię i nazwisko", "Data", "Grupa"]
    for i, (cell, width, label) in enumerate(zip(table.rows[0].cells, widths, labels)):
        cell.width = Cm(width)
        set_cell_shading(cell, PALE_GRAY)
        set_cell_margins(cell, top=85, bottom=150)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_paragraph(p, after=0)
        add_text(p, label + "\n", size=8.5, bold=True, color="4B5563")
        add_text(p, " " * 10, size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_strategy(doc):
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=5)
    add_text(p, "Jak pracować: ", size=9.5, bold=True, color=BLUE)
    add_text(p, "najpierw przeczytaj cały tekst, potem wracaj do niego przy każdym pytaniu. Podkreślaj słowa, które pomagają znaleźć odpowiedź.", size=9.5)


def add_section(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    add_text(p, title, size=13, bold=True, color="000000")
    return p


def add_subsection(doc, title):
    p = doc.add_paragraph(style="Heading 2")
    add_text(p, title, size=11.5, bold=True, color="000000")
    return p


def add_reading_text(doc, paragraphs):
    for text in paragraphs:
        p = doc.add_paragraph()
        set_paragraph(p, after=5, line=1.08)
        p.paragraph_format.first_line_indent = Cm(0.55)
        add_text(p, text, size=10.5)


def add_self_check(p):
    """Przy sprawdzaniu uczen sam zakresla + albo - przy zadaniu (prawy margines)."""
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.9), WD_TAB_ALIGNMENT.RIGHT)
    add_text(p, "	( + )  ( − )", size=10, bold=True, color="6B7280")


def add_task(doc, number, prompt, lines=0, options=None, compact=False):
    p = doc.add_paragraph()
    set_paragraph(p, before=2 if not compact else 0, after=2, line=1.03, keep=bool(options))
    if number:
        add_text(p, f"{number}. ", size=10.3, bold=True, color=BLUE)
    add_text(p, prompt, size=10.3)
    if number:
        add_self_check(p)
    if options:
        for i, option in enumerate(options):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm(0.55)
            set_paragraph(op, after=7 if i == len(options) - 1 else 1.5, line=1.0)
            add_text(op, "☐ ", size=13)
            add_text(op, option, size=10.1)
    for _ in range(lines):
        lp = doc.add_paragraph()
        set_paragraph(lp, after=2, line=1.0)
        add_text(lp, "................................................................................................................................", size=9, color="9CA3AF")


def add_two_col_table(doc, rows, left_width=8.7, right_width=8.7, header=None):
    total_rows = len(rows) + (1 if header else 0)
    table = doc.add_table(rows=total_rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    start = 0
    if header:
        for i, txt in enumerate(header):
            cell = table.rows[0].cells[i]
            set_cell_shading(cell, BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell, top=100, bottom=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph(p, after=0)
            add_text(p, txt, size=9.5, bold=True, color="FFFFFF")
        set_repeat_table_header(table.rows[0])
        start = 1
    for r_idx, row in enumerate(rows, start=start):
        for c_idx, txt in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.width = Cm(left_width if c_idx == 0 else right_width)
            set_cell_borders(cell)
            set_cell_margins(cell, top=90, bottom=90)
            if (r_idx - start) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            set_paragraph(p, after=0, line=1.0)
            add_text(p, txt, size=9.8)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_page_two_header(doc, number, title):
    p = doc.add_paragraph()
    set_paragraph(p, after=5)
    add_text(p, f"Arkusz {number}  {title}  |  strona 2", size=9, bold=True, color="4B5563")


def add_after_check(doc):
    """Koniec arkusza robimy razem po sprawdzeniu - w osobnej ramce, zeby nikt nie robil go w trakcie."""
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.rows[0].cells[0]
    cell.width = Cm(17.6)
    set_cell_shading(cell, "F4F8FB")
    set_cell_borders(cell, color=BLUE, size="10")
    for edge in cell._tc.tcPr.find(qn("w:tcBorders")):
        edge.set(qn("w:val"), "dashed")
    set_cell_margins(cell, top=120, start=180, bottom=140, end=180)

    p = cell.paragraphs[0]
    set_paragraph(p, after=4)
    add_text(p, "Po wspólnym sprawdzeniu ", size=11.5, bold=True, color=BLUE)
    add_text(p, "- tę część robimy razem na końcu lekcji", size=9.5, color="4B5563")

    p = cell.add_paragraph()
    set_paragraph(p, after=2)
    add_text(p, "1. Policz swoje plusy. ", size=10, bold=True)
    add_text(p, "Przy sprawdzaniu zakreśl obok każdego zadania + (dobrze) albo − (do poprawy). Policz plusy i wpisz wynik.", size=10)

    p = cell.add_paragraph()
    set_paragraph(p, before=2, after=5)
    add_text(p, "Moje plusy: ", size=9.8)
    add_text(p, "..............", size=9.8, color="9CA3AF")
    add_text(p, " / 12          ", size=9.8)
    for label, plus in (("4-8 = ", "+"), ("     9-10 = ", "++"), ("     11-12 = ", "+++")):
        add_text(p, label, size=9.8)
        add_text(p, plus, size=9.8, bold=True)

    p = cell.add_paragraph()
    set_paragraph(p, after=2)
    add_text(p, "2. Zamaluj numer zadania, które teraz rozumiem lepiej.", size=10, bold=True)
    circles = cell.add_table(rows=1, cols=12)
    circles.autofit = False
    for n in range(12):
        c = circles.rows[0].cells[n]
        c.width = Cm(1.05)
        set_cell_borders(c, color="374151")
        set_cell_shading(c, "FFFFFF")
        set_cell_margins(c, top=90, start=40, bottom=90, end=40)
        cp = c.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(cp, after=0, line=1.0)
        add_text(cp, str(n + 1), size=10, bold=True)
    cell.add_paragraph().paragraph_format.space_after = Pt(0)


WORKSHEETS = [
    {
        "number": 1,
        "title": "Kiermasz dobrych pomysłów",
        "focus": "Czytanie ze zrozumieniem  kolejność wydarzeń  rzeczownik  zdanie  krótka wiadomość",
        "text": [
            "W piątek klasa IV a miała przygotować stoisko na szkolny kiermasz. Pani poprosiła uczniów, aby przynieśli drobiazgi wykonane własnoręcznie. Lena zrobiła zakładki do książek, Olek ulepił małe figurki z masy solnej, a Sara zaproponowała, że przygotuje karteczki z cenami.",
            "Po lekcjach dzieci spotkały się w sali. Szybko zauważyły, że na jednej ławce nie zmieszczą się wszystkie przedmioty. Olek chciał schować część rzeczy do pudełka, lecz Lena miała inny pomysł. Uczniowie połączyli dwie ławki, przykryli je jasnym obrusem i ustawili wyroby grupami. Najmniejsze rzeczy położyli z przodu, żeby każdy mógł je zobaczyć.",
            "Następnego dnia ich stoisko odwiedziło wiele osób. Najpierw sprzedały się zakładki, a potem figurki. Pod koniec kiermaszu została tylko jedna kartka z napisem: Dziękujemy za pomoc. Uczniowie byli zmęczeni, ale dumni, bo wspólnie rozwiązali problem i zebrali pieniądze na nowe książki do biblioteki.",
        ],
        "reading": [
            ("1", "Po co klasa przygotowała stoisko?", 1, None),
            ("2", "Kto przygotował karteczki z cenami?", 0, ["Lena", "Olek", "Sara"]),
            ("3", "Ponumeruj wydarzenia od 1 do 4.", 0, None),
            ("", "□ Uczniowie połączyli dwie ławki.   □ Sprzedały się zakładki.\n□ Dzieci przyniosły własne wyroby.   □ Klasa zebrała pieniądze na książki.", 0, None),
            ("4", "Dlaczego najmniejsze rzeczy ustawiono z przodu? Odpowiedz pełnym zdaniem.", 1, None),
            ("5", "Które zdanie najlepiej mówi, czego uczy ta historia?", 0, ["Warto działać wspólnie, gdy pojawia się problem.", "Na kiermaszu należy sprzedawać tylko książki.", "Najlepiej pracuje się bez planu."]),
            ("6", "Znajdź w tekście słowo, które znaczy prawie to samo co przedmioty niewielkiej wartości.", 1, None),
        ],
        "language": [
            ("7", "Wypisz z pierwszego akapitu trzy rzeczowniki.", 1, None),
            ("8", "Uzupełnij zdanie właściwą formą wyrazu książka: Uczniowie zebrali pieniądze na nowe ....................... do biblioteki.", 0, None),
            ("9", "Ułóż wyrazy w kolejności alfabetycznej: stoisko, zakładka, obrus, figurka.", 1, None),
            ("10", "Połącz dwa zdania wyrazem ponieważ: Uczniowie byli dumni. Wspólnie rozwiązali problem.", 1, None),
            ("11", "Popraw zdanie. Zapisz wielką literę i potrzebny znak na końcu: czy sara przygotowała ceny", 1, None),
        ],
        "writing": "Napisz krótką wiadomość do kolegi lub koleżanki. Zaproś tę osobę na szkolny kiermasz. Podaj miejsce, dzień i jedną zachętę.",
        "writing_lines": 8,
    },
    {
        "number": 2,
        "title": "Klucz z numerem siedem",
        "focus": "Czytanie ze zrozumieniem  wnioskowanie  części mowy  ortografia  opis przedmiotu",
        "text": [
            "Maja znalazła na szkolnym korytarzu mały mosiężny klucz. Leżał tuż obok schodów i miał przyczepioną okrągłą zawieszkę z numerem siedem. Dziewczynka podniosła go i od razu zaniosła do sekretariatu. Sekretarka obejrzała zgubę, ale nie wiedziała, do czego pasuje.",
            "Na długiej przerwie Maja opowiedziała o znalezisku Kubie. Chłopiec przypomniał sobie, że w sali przyrodniczej stoi stara szafka z ponumerowanymi szufladami. Poszli tam razem z nauczycielką. Klucz rzeczywiście otworzył siódmą szufladę. W środku nie było skarbu, lecz pudełko z lupami, miarkami i kompasami.",
            "Nauczycielka ucieszyła się, bo od kilku dni szukała zapasowego klucza. Podziękowała Mai za uczciwość, a Kubie za dobrą pamięć. Na następnej lekcji cała klasa wykorzystała odnalezione przyrządy podczas obserwacji kamieni. Maja uznała, że zwykły klucz może rozpocząć całkiem ciekawą przygodę.",
        ],
        "reading": [
            ("1", "Gdzie Maja znalazła klucz?", 1, None),
            ("2", "Co znajdowało się w siódmej szufladzie? Wymień dwa przedmioty.", 1, None),
            ("3", "Dlaczego Kuba pomógł rozwiązać zagadkę klucza?", 0, ["Miał dobrą pamięć i przypomniał sobie o szafce.", "Znalazł drugi klucz na boisku.", "Pracował w sekretariacie."]),
            ("4", "Zaznacz P, jeśli zdanie jest prawdziwe, albo F, jeśli jest fałszywe.", 0, None),
            ("", "Klucz miał kwadratową zawieszkę.  P / F\nNauczycielka szukała zapasowego klucza.  P / F\nUczniowie wykorzystali przyrządy na lekcji.  P / F", 0, None),
            ("5", "Jaką cechę Mai pokazuje jej zachowanie? Podaj cechę i dowód z tekstu.", 2, None),
            ("6", "Nadaj tekstowi inny tytuł. Tytuł powinien pasować do całej historii.", 1, None),
        ],
        "language": [
            ("7", "W zdaniu Maja znalazła mały klucz podkreśl rzeczownik, otocz czasownik i zaznacz przymiotnik falistą linią.", 0, None),
            ("8", "Zapisz liczbę mnogą: klucz - .......................   szuflada - .......................   kamień - .......................", 0, None),
            ("9", "Wybierz poprawny zapis: drużyna / drurzyna, korytarz / korytasz, mosiężny / mosienżny.", 1, None),
            ("10", "Utwórz rodzinę wyrazu klucz. Dopisz dwa wyrazy pokrewne.", 1, None),
            ("11", "Rozwiń zdanie Maja znalazła klucz, dodając odpowiedź na pytania gdzie i kiedy.", 1, None),
        ],
        "writing": "Opisz znaleziony klucz w 3-4 zdaniach. Napisz, jak wyglądał, gdzie leżał i co zrobiła osoba, która go znalazła.",
        "writing_lines": 8,
    },
    {
        "number": 3,
        "title": "Sen pomaga się uczyć",
        "focus": "Tekst informacyjny  przyczyna i skutek  fakt i opinia  czasownik  notatka",
        "text": [
            "Sen nie jest przerwą w pracy organizmu. Kiedy śpimy, mózg porządkuje informacje z całego dnia. Dzięki temu łatwiej zapamiętujemy nowe wiadomości i sprawniej rozwiązujemy zadania. Odpoczynek pomaga też zachować uwagę podczas lekcji.",
            "Dzieci w wieku szkolnym zwykle potrzebują więcej snu niż dorośli. Jeśli ktoś kładzie się bardzo późno, rano może być senny, rozdrażniony i mniej skupiony. Jedna długa noc nie naprawi od razu kilku dni zbyt krótkiego snu. Najlepiej odpoczywać regularnie.",
            "Dobrym przygotowaniem do snu jest spokojny wieczór. Warto odłożyć telefon lub tablet, przewietrzyć pokój i przygotować rzeczy na następny dzień. Pomaga także podobna pora zasypiania. Każdy może sprawdzić, po jakim wieczornym zwyczaju budzi się bardziej wypoczęty.",
        ],
        "reading": [
            ("1", "Co mózg robi z informacjami podczas snu?", 1, None),
            ("2", "Podaj dwie trudności, które mogą pojawić się po zbyt krótkim śnie.", 1, None),
            ("3", "Zaznacz trzy czynności, które pomagają przygotować się do snu.", 0, ["przewietrzenie pokoju", "długa gra na tablecie", "przygotowanie rzeczy na rano", "podobna pora zasypiania"]),
            ("4", "Uzupełnij związek przyczyny i skutku: Ktoś kładzie się bardzo późno, dlatego rano ...................................", 1, None),
            ("5", "Które zdanie jest główną myślą tekstu?", 0, ["Regularny sen pomaga w nauce i codziennym funkcjonowaniu.", "Każdy powinien spać dokładnie tyle samo godzin.", "Wieczorem trzeba przygotować tylko tornister."]),
            ("6", "Zdanie Spokojny wieczór jest najlepszą częścią dnia to fakt czy opinia? Uzasadnij krótko.", 1, None),
        ],
        "language": [
            ("7", "Wypisz z drugiego akapitu trzy czasowniki.", 1, None),
            ("8", "Zmień zdanie na czas przeszły: Mózg porządkuje informacje.", 1, None),
            ("9", "Wybierz poprawny zapis: mózg / musk, późno / puźno, przygotować / pszygotować.", 1, None),
            ("10", "Ułóż jedno zdanie oznajmujące i jedno pytające o odpoczynku.", 2, None),
            ("11", "Podziel wyraz wypoczęty na sylaby i zaznacz samogłoski.", 1, None),
        ],
        "writing": "Napisz notatkę punktową pod tytułem Mój sposób na dobry wieczór. Wypisz trzy czynności w sensownej kolejności - każdą w osobnym punkcie.",
        "writing_lines": 8,
    },
    {
        "number": 4,
        "title": "Droga do starego dębu",
        "focus": "Instrukcja  plan działania  kierunki  przymiotnik  ortografia  opis trasy",
        "text": [
            "Klasa IV wybrała się do parku na grę terenową. Każda drużyna otrzymała mapę i krótką instrukcję. Celem było odnalezienie starego dębu, a przy nim koperty z ostatnim zadaniem. Uczniowie mieli poruszać się razem i zaznaczać na mapie wszystkie mijane miejsca.",
            "Instrukcja drużyny Mai brzmiała: Od bramy idź prosto do fontanny. Skręć w lewo i przejdź obok dwóch ławek. Za małym mostkiem wybierz ścieżkę po prawej stronie. Idź nią aż do tablicy z rysunkiem wiewiórki. Stary dąb rośnie dziesięć kroków za tablicą.",
            "Przy mostku Tomek chciał skręcić w lewo, bo zobaczył szeroką alejkę. Maja ponownie przeczytała instrukcję i wskazała właściwą drogę. Drużyna dotarła do dębu bez błądzenia. W kopercie znalazła pytanie: Co było ważniejsze - szybki marsz czy uważne czytanie? Wszyscy zgodnie zapisali drugą odpowiedź.",
        ],
        "reading": [
            ("1", "Jaki był cel gry terenowej?", 1, None),
            ("2", "Co uczniowie mieli zaznaczać na mapie?", 1, None),
            ("3", "Wpisz brakujące miejsca w planie drogi.", 0, None),
            ("", "brama → ........................ → dwie ławki → ........................ → tablica z wiewiórką → stary dąb", 1, None),
            ("4", "Dlaczego Maja przeczytała instrukcję ponownie?", 1, None),
            ("5", "Która umiejętność pomogła drużynie dotrzeć do celu? Podaj odpowiedź i dowód z tekstu.", 2, None),
            ("6", "Wyjaśnij własnymi słowami wyrażenie dotarła bez błądzenia.", 1, None),
        ],
        "language": [
            ("7", "Dopisz po jednym przymiotniku do rzeczowników.\ndąb - ..................   alejka - ..................   instrukcja - ..................", 0, None),
            ("8", "Ułóż w kolejności alfabetycznej: brama, mostek, dąb, fontanna.", 1, None),
            ("9", "Wybierz poprawny zapis: drużyna / drurzyna, ścieżka / ścieszka, wiewiórka / wiewjurka.", 1, None),
            ("10", "Zmień zdanie Idź prosto do fontanny na zdanie grzecznej prośby.", 1, None),
            ("11", "Rozwiń zdanie Drużyna dotarła, dodając co najmniej cztery wyrazy.", 1, None),
        ],
        "writing": "Napisz instrukcję dojścia od drzwi sali do wybranego miejsca w szkole. Użyj co najmniej trzech określeń kierunku, na przykład prosto, w lewo, obok.",
        "writing_lines": 8,
    },
]


LEVEL_UP = "Poziom wyżej"
PFQ = "Zaznacz P (prawda), F (fałsz) albo ? (tekst o tym nie mówi)."

WORKSHEETS += [
    {
        "number": 5,
        "level": LEVEL_UP,
        "title": "Drugie miejsce",
        "focus": "Wnioskowanie o uczuciach  intencje bohaterów  odmiana rzeczownika  przecinek  dialog",
        "text": [
            "Na tablicy przy wejściu wisiała lista nagrodzonych w szkolnym konkursie na komiks. Antek przeczytał ją dwa razy, zanim zauważył swoje nazwisko. Było na drugim miejscu, tuż pod nazwiskiem Igi z równoległej klasy. Antek wrócił do sali i przez całą lekcję matematyki rysował w zeszycie kwadraty, chociaż pani tłumaczyła nowe zadania.",
            "Na przerwie Iga podeszła do niego z teczką pod pachą.",
            "– Widziałam twój komiks o kocie astronaucie – powiedziała. – Mnie nigdy nie wychodzi ruch. A u ciebie ten kot naprawdę leci.",
            "Antek wzruszył ramionami i mruknął, że jury i tak wybrało jej pracę. Iga przez chwilę milczała. Potem otworzyła teczkę i pokazała mu kartkę pełną przekreślonych szkiców.",
            "– Ten smok wyszedł mi dopiero za szóstym razem – przyznała.",
            "Po lekcjach Antek nie poszedł od razu do domu. Usiadł na ławce przed szkołą, wyjął swój komiks i długo przyglądał się ostatniemu obrazkowi. Wieczorem narysował coś nowego. Następnego ranka na ławce Igi leżała kartka: zielony smok leciał obok kota w skafandrze, a nad nimi unosił się dymek z pytaniem: Zrobimy razem drugi odcinek?",
        ],
        "reading": [
            ("1", "Kto zajął w konkursie pierwsze miejsce, a kto drugie?", 1, None),
            ("2", "Antek przez całą lekcję rysował kwadraty. Co to mówi o tym, jak się wtedy czuł?", 1, None),
            ("3", "Po co Iga pokazała Antkowi przekreślone szkice?", 0, ["Chciała się pochwalić, ile już narysowała.", "Chciała pokazać, że jej też nie wszystko wychodzi od razu.", "Chciała, żeby Antek poprawił jej rysunki."]),
            ("4", PFQ, 0, None),
            ("", "Antek i Iga chodzą do tej samej klasy.  P / F / ?\nIga rysuje komiksy dłużej niż Antek.  P / F / ?\nIga podziwia, jak Antek rysuje ruch.  P / F / ?", 0, None),
            ("5", "Wyjaśnij własnymi słowami, co Iga chciała powiedzieć zdaniem: U ciebie ten kot naprawdę leci.", 1, None),
            ("6", "Jak zmienia się nastrój Antka od początku do końca historii? Podaj dowód z tekstu.", 2, None),
        ],
        "language": [
            ("7", "Wpisz wyraz komiks w odpowiedniej formie.\nNa liście nie było jego ......................... na pierwszym miejscu.\nAntek długo przyglądał się swojemu .........................  Iga zachwyciła się jego .........................\nNastępnego dnia rozmawiali o nowym .........................", 0, None),
            ("8", "Dopisz wyraz, który wyjaśnia pisownię ó albo rz.\nwróciła - .........................   szkółka - .........................   w kolorze - .........................   malarz - .........................", 0, None),
            ("9", "Wstaw brakujące przecinki (są trzy): Antek wiedział że jego kot dobrze wyszedł ale wciąż było mu przykro bo marzył o pierwszym miejscu.", 0, None),
            ("10", "Jak zmieni się sens zdania, jeśli zamiast Antek mruknął napiszemy Antek krzyknął?", 1, None),
            ("11", "Napisz jedno zdanie o Antku albo o Idze. Połącz w nim dwie myśli wyrazem ale albo dlatego.", 1, None),
        ],
        "writing": "Napisz dalszy ciąg historii: co powiedziała Iga, gdy znalazła kartkę? Napisz 5-6 zdań. Użyj dialogu - każdą wypowiedź zacznij od nowej linijki i od myślnika.",
        "writing_lines": 8,
    },
    {
        "number": 6,
        "level": LEVEL_UP,
        "title": "Dwie drogi do Afryki",
        "focus": "Tekst popularnonaukowy  porównanie  łączenie informacji  słowotwórstwo  ortografia",
        "text": [
            "Każdej jesieni nad Polską odbywa się wielka wędrówka. Ptaki, które latem karmiły pisklęta na naszych łąkach, odlatują do Afryki. Nie robią tego z powodu zimna, jak często się sądzi. Wiele z nich zniosłoby niską temperaturę, ale zimą brakuje im pokarmu: owadów, żab czy dżdżownic.",
            "Bocian biały to ciężki ptak o szerokich skrzydłach. Gdyby przez całą drogę nimi machał, szybko by się zmęczył. Dlatego szybuje na prądach ciepłego powietrza, które unoszą się nad nagrzaną ziemią. Nad morzem takich prądów prawie nie ma, więc bociany lecą dłuższą drogą nad lądem: z Polski przez Turcję i Bliski Wschód. Wyruszają w sierpniu, a młode ptaki często startują przed dorosłymi i same muszą odnaleźć kierunek.",
            "Jaskółka dymówka jest lekka i zwinna. Leci, machając skrzydłami, a po drodze łapie owady w locie, więc nie musi robić długich przerw na posiłek. Przelatuje nad Morzem Śródziemnym i nad pustynią. Zanim ruszy na południe, zbiera się z innymi jaskółkami na przewodach elektrycznych.",
            "Naukowcy wciąż badają, jak ptaki nie gubią drogi. Wiadomo, że pomagają im słońce, gwiazdy i ukształtowanie terenu, a niektóre gatunki wyczuwają pole magnetyczne Ziemi, tak jak igła kompasu. Mimo to wędrówka jest niebezpieczna i nie wszystkie ptaki wracają wiosną do swoich gniazd.",
        ],
        "reading": [
            ("1", "Z jakiego powodu ptaki naprawdę odlatują? Jakie błędne przekonanie prostuje autor?", 1, None),
            ("2", "Porównaj oba ptaki. Uzupełnij tabelę.", 0, {"table": (
                ["", "bocian biały", "jaskółka dymówka"],
                [["budowa", "", ""], ["sposób lotu", "", ""], ["lot nad morzem", "", ""]],
            )}),
            ("3", "Dlaczego bociany nie lecą nad morzem? Połącz dwie informacje z tekstu.", 1, None),
            ("4", PFQ, 0, None),
            ("", "Bociany odlatują, bo nie znoszą zimna.  P / F / ?\nMłode bociany zawsze lecą razem z rodzicami.  P / F / ?\nJaskółki żyją dłużej niż bociany.  P / F / ?", 0, None),
            ("5", "Po co autor napisał, że ptaki wyczuwają pole magnetyczne tak jak igła kompasu?", 0, ["Żeby wyjaśnić trudną rzecz czymś, co czytelnik zna.", "Żeby pokazać, że ptaki noszą ze sobą kompasy.", "Żeby opisać, jak wygląda bocian."]),
            ("6", "Podaj dwa zagrożenia w czasie wędrówki, o których można wnioskować z tekstu.", 1, None),
        ],
        "language": [
            ("7", "Utwórz przymiotniki od rzeczowników. Uważaj na pisownię.\nzima - ....................   ptak - ....................   Afryka - ....................   morze - ....................   słońce - ....................", 0, None),
            ("8", "Wpisz ż, rz albo sz:  ..aby   d..d..ownice   pr..ewody   wyru..ają   mo..e   ..ółty dziób", 0, None),
            ("9", "Wpisz wyraz ptak w odpowiedniej formie.\nNad łąką krążyło pięć .........................  Naukowcy przyglądają się ......................... przez lornetkę.\nW książce przeczytałem o ......................... wędrownych.", 0, None),
            ("10", "Połącz zdania na dwa sposoby: raz wyrazem więc, raz wyrazem ponieważ. Zimą brakuje owadów. Jaskółki odlatują do Afryki.", 2, None),
            ("11", "W tekście igła to igła kompasu. Ułóż dwa zdania, w których wyraz igła ma dwa inne znaczenia.", 2, None),
        ],
        "writing": "Wyjaśnij młodszemu koledze albo koleżance, dlaczego bocian leci do Afryki okrężną drogą. Napisz 4-5 zdań własnymi słowami - nie przepisuj zdań z tekstu.",
        "writing_lines": 7,
    },
    {
        "number": 7,
        "level": LEVEL_UP,
        "title": "Ktoś nowy w domu",
        "focus": "Kto opowiada?  wnioskowanie  czasownik  ę czy ą  opowiadanie z innej perspektywy",
        "text": [
            "Od trzech miesięcy w naszym mieszkaniu mieszka ktoś nowy. Jest mały, głośny i pachnie mlekiem. Kiedy płacze, wszyscy biegną do niego, nawet jeśli właśnie trzymam w zębach piłkę i czekam przy drzwiach. Wcześniej to ja byłem najważniejszy. Pani mówiła do mnie cienkim głosem, a pan codziennie rano zabierał mnie do parku, zanim wypił kawę.",
            "Teraz spacery są krótsze. Pan ziewa, schodzi po schodach w kapciach i wraca, zanim zdążę obwąchać wszystkie krzaki. Wczoraj zjadłem skarpetkę z małego kosza. Nikt na mnie nie krzyczał. Pani tylko westchnęła i powiedziała, że chyba czuję się samotny. Nie wiem, co znaczy samotny, ale dodatkowy kawałek kiełbasy w misce bardzo mi się spodobał.",
            "Dziś rano stało się coś dziwnego. Nowy leżał na kocu na podłodze i patrzył prosto na mnie. Podszedłem ostrożnie, a on wyciągnął rękę i złapał mnie za ucho. Nie bolało. Pani zamarła, ale ja się nie ruszyłem. Wtedy Nowy zaśmiał się pierwszy raz w życiu. Pan powiedział, że to najlepszy prezent. Nie rozumiem ludzi - przecież niczego mu nie dałem. Może jednak zostanę jego starszym bratem.",
        ],
        "reading": [
            ("1", "Kto opowiada tę historię? Podaj trzy dowody z tekstu.", 2, None),
            ("2", "Kim jest Nowy? Po czym to poznajesz?", 1, None),
            ("3", "Dlaczego narrator zjadł skarpetkę? Czy pani dobrze odgadła, co czuje? Uzasadnij.", 1, None),
            ("4", PFQ, 0, None),
            ("", "Nowy mieszka w domu od trzech dni.  P / F / ?\nPani nakrzyczała na narratora za skarpetkę.  P / F / ?\nNarrator ma na imię Burek.  P / F / ?", 0, None),
            ("5", "Pan powiedział, że to najlepszy prezent. Wyjaśnij, co miał na myśli.", 1, None),
            ("6", "Jak zmienia się stosunek narratora do Nowego? Co znaczy ostatnie zdanie?", 1, None),
        ],
        "language": [
            ("7", "Opowiedz te zdania tak, jakby mówił je ktoś z boku. Zmień formy czasowników: Wcześniej to ja byłem najważniejszy. Wczoraj zjadłem skarpetkę.", 1, None),
            ("8", "Określ osobę, liczbę i czas.\nbiegną - ...........................................   zjadłem - ...........................................   zostanę - ...........................................", 0, None),
            ("9", "Co znaczy w tekście zdanie Pani zamarła?", 0, ["Pani bardzo zmarzła.", "Pani znieruchomiała z wrażenia.", "Pani nagle zasnęła."]),
            ("10", "Wpisz ę albo ą: Pies trzyma w zębach piłk.. .  Pan rzuca piłk.. daleko.  Pies bawi się piłk.. .  Pani przykryła dziecko kołdr.. .\nKiedy piszesz ę, a kiedy ą? Pomóż sobie pytaniami kogo? co? oraz kim? czym?", 1, None),
            ("11", "Jakie uczucie miał narrator, gdy wszyscy biegli do płaczącego dziecka? Nazwij je i dopisz dwa wyrazy z nim spokrewnione.", 1, None),
        ],
        "writing": "Opisz jeden dzień z życia rodziny oczami zwierzęcia, które mieszka w domu. Nie zdradzaj, kim jest narrator - czytelnik ma to odgadnąć. Napisz 5-6 zdań.",
        "writing_lines": 8,
    },
]


def add_grid(doc, header, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.6] + [6.9] * (len(header) - 1)
    for r_idx, row in enumerate([header] + rows):
        for c_idx, txt in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.width = Cm(widths[c_idx])
            set_cell_borders(cell)
            set_cell_margins(cell, top=110, bottom=110)
            if r_idx == 0 and txt:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            set_paragraph(p, after=0, line=1.0)
            add_text(p, txt, size=9.8, bold=r_idx == 0 or c_idx == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def slug(title):
    table = str.maketrans("łóąęśżźćń", "loaeszzcn")
    return title.lower().translate(table).replace(" ", "-")


def file_stem(ws):
    return f"arkusz-{ws['number']}-{slug(ws['title'])}"


def build_worksheet(ws):
    doc = Document()
    configure_doc(doc)
    level = ws.get("level")
    add_footer(doc, f"Arkusz {ws['number']}" + (f"  |  {level.lower()}" if level else ""))
    add_identity(doc, ws["number"], ws["title"], (f"{level}  |  " if level else "") + ws["focus"])
    if level:
        p = doc.add_paragraph()
        set_paragraph(p, before=0, after=5)
        add_text(p, "Jak pracować: ", size=9.5, bold=True, color=BLUE)
        add_text(p, "odpowiedzi nie zawsze są zapisane wprost. Szukaj wskazówek, łącz informacje z różnych akapitów i sprawdzaj, czy tekst naprawdę o czymś mówi.", size=9.5)
    else:
        add_strategy(doc)
    add_section(doc, "Przeczytaj tekst")
    add_reading_text(doc, ws["text"])
    add_section(doc, "Sprawdź, co rozumiesz")
    for n, prompt, lines, options in ws["reading"]:
        if isinstance(options, dict):
            add_task(doc, n, prompt, compact=True)
            add_grid(doc, *options["table"])
        else:
            add_task(doc, n, prompt, lines=lines, options=options, compact=True)

    add_page_break(doc)
    add_page_two_header(doc, ws["number"], ws["title"])
    add_section(doc, "Poćwicz język")
    for n, prompt, lines, options in ws["language"]:
        add_task(doc, n, prompt, lines=lines, options=options)

    add_section(doc, "Napisz krótko")
    p = doc.add_paragraph()
    set_paragraph(p, after=2)
    add_text(p, "12. ", size=10.3, bold=True, color=BLUE)
    add_text(p, ws["writing"], size=10.3)
    add_self_check(p)
    for _ in range(ws["writing_lines"]):
        p = doc.add_paragraph()
        set_paragraph(p, after=2, line=1.0)
        add_text(p, "................................................................................................................................", size=9, color="9CA3AF")

    add_after_check(doc)

    path = OUT / f"{file_stem(ws)}.docx"
    doc.save(path)
    return path


ANSWERS = {
    1: [
        ("1", "Klasa przygotowała stoisko na szkolny kiermasz, aby sprzedać własnoręcznie wykonane przedmioty i zebrać pieniądze na książki do biblioteki."),
        ("2", "Sara."),
        ("3", "2, 3, 1, 4 - w kolejności pozycji zapisanych w zadaniu."),
        ("4", "Ustawiono je z przodu, żeby każdy mógł je zobaczyć."),
        ("5", "Warto działać wspólnie, gdy pojawia się problem."),
        ("6", "drobiazgi"),
        ("7", "Przykładowo: klasa, stoisko, kiermasz, pani, uczniowie, drobiazgi, Lena, zakładki, książki, Olek, figurki, masa, Sara, karteczki, ceny."),
        ("8", "książki"),
        ("9", "figurka, obrus, stoisko, zakładka"),
        ("10", "Uczniowie byli dumni, ponieważ wspólnie rozwiązali problem."),
        ("11", "Czy Sara przygotowała ceny?"),
        ("12", "Wypowiedź zawiera zaproszenie, miejsce, dzień i zachętę. Jest zrozumiała i poprawnie zakończona."),
    ],
    2: [
        ("1", "Na szkolnym korytarzu, obok schodów."),
        ("2", "Dwa spośród: lupy, miarki, kompasy."),
        ("3", "Miał dobrą pamięć i przypomniał sobie o szafce."),
        ("4", "F, P, P."),
        ("5", "Na przykład: uczciwa - zaniosła znaleziony klucz do sekretariatu."),
        ("6", "Tytuł powinien dotyczyć całej historii, na przykład Zagadka siódmej szuflady."),
        ("7", "rzeczowniki: Maja, klucz; czasownik: znalazła; przymiotnik: mały."),
        ("8", "klucze, szuflady, kamienie"),
        ("9", "drużyna, korytarz, mosiężny"),
        ("10", "Przykładowo: kluczyk, kluczowy. Należy uznać inne poprawne wyrazy pokrewne."),
        ("11", "Przykładowo: Wczoraj rano Maja znalazła mały klucz na szkolnym korytarzu."),
        ("12", "Opis obejmuje wygląd, miejsce znalezienia i zachowanie znalazcy. Ma 3-4 zdania."),
    ],
    3: [
        ("1", "Porządkuje informacje z całego dnia."),
        ("2", "Dwie spośród: senność, rozdrażnienie, mniejsze skupienie."),
        ("3", "przewietrzenie pokoju; przygotowanie rzeczy na rano; podobna pora zasypiania"),
        ("4", "Na przykład: może być senny, rozdrażniony i mniej skupiony."),
        ("5", "Regularny sen pomaga w nauce i codziennym funkcjonowaniu."),
        ("6", "Opinia, ponieważ wyraża osobistą ocenę i nie można jej sprawdzić tak samo u wszystkich."),
        ("7", "Przykładowo: potrzebują, kładzie się, może, być, naprawi, odpoczywać."),
        ("8", "Mózg porządkował informacje."),
        ("9", "mózg, późno, przygotować"),
        ("10", "Dowolne dwa poprawne zdania zgodne z wymaganym typem."),
        ("11", "wy-po-czę-ty; samogłoski: y, o, ę, y"),
        ("12", "Notatka ma tytuł i trzy ponumerowane punkty w sensownej kolejności. Każdy punkt to krótkie hasło, np. 1. Odkładam telefon. 2. Pakuję plecak na jutro. 3. Wietrzę pokój i kładę się o stałej porze."),
    ],
    4: [
        ("1", "Odnalezienie starego dębu i koperty z ostatnim zadaniem."),
        ("2", "Wszystkie mijane miejsca."),
        ("3", "fontanna; mały mostek"),
        ("4", "Tomek chciał wybrać niewłaściwą drogę, więc Maja sprawdziła instrukcję."),
        ("5", "Uważne czytanie - Maja ponownie przeczytała instrukcję i wskazała właściwą drogę."),
        ("6", "Drużyna dotarła do celu, nie gubiąc drogi i nie wybierając złych ścieżek."),
        ("7", "Dowolne poprawne połączenia, na przykład: stary dąb, szeroka alejka, krótka instrukcja."),
        ("8", "brama, dąb, fontanna, mostek"),
        ("9", "drużyna, ścieżka, wiewiórka"),
        ("10", "Na przykład: Proszę, idź prosto do fontanny."),
        ("11", "Dowolne poprawne rozwinięcie o co najmniej cztery wyrazy, na przykład: Drużyna szybko dotarła do starego dębu w parku."),
        ("12", "Instrukcja prowadzi od drzwi sali do wybranego miejsca i zawiera co najmniej trzy określenia kierunku."),
    ],
    5: [
        ("1", "Pierwsze miejsce zajęła Iga, drugie Antek."),
        ("2", "Był rozczarowany, smutny, myślami przy konkursie - nie słuchał lekcji, tylko bezmyślnie rysował."),
        ("3", "Chciała pokazać, że jej też nie wszystko wychodzi od razu."),
        ("4", "F (Iga jest z równoległej klasy), ?, P. Omówić różnicę między F a ?: fałsz to coś sprzecznego z tekstem, ? to coś, o czym tekst nie mówi."),
        ("5", "Na przykład: narysowałeś kota tak, że widać, jak się porusza / rysunek wygląda jak prawdziwy lot. Kot nie leci naprawdę - to pochwała rysunku."),
        ("6", "Na początku rozczarowany i zazdrosny (rysuje kwadraty, mruczy, że jury wybrało jej pracę), na końcu pogodzony, chętny do współpracy (rysuje smoka z kotem i proponuje wspólny odcinek)."),
        ("7", "komiksu, komiksowi, komiksem, komiksie"),
        ("8", "wracać; szkoła; kolor; malarski (lub malarka)"),
        ("9", "Antek wiedział, że jego kot dobrze wyszedł, ale wciąż było mu przykro, bo marzył o pierwszym miejscu."),
        ("10", "Krzyknął znaczy głośno, ze złością albo z emocjami; mruknął - cicho, niechętnie, pod nosem. Zmienia się obraz zachowania Antka."),
        ("11", "Dowolne poprawne zdanie z przecinkiem przed ale albo dlatego, np. Iga wygrała, ale doceniła rysunek Antka."),
        ("12", "Dalszy ciąg pasuje do historii, ma 5-6 zdań i co najmniej dwie wypowiedzi zapisane od nowej linijki z myślnikiem."),
    ],
    6: [
        ("1", "Zimą brakuje im pokarmu (owadów, żab, dżdżownic). Autor prostuje przekonanie, że ptaki odlatują z powodu zimna."),
        ("2", "budowa: bocian ciężki, szerokie skrzydła / jaskółka lekka, zwinna; lot: bocian szybuje na prądach ciepłego powietrza / jaskółka macha skrzydłami; nad morzem: bocian omija morze / jaskółka przelatuje nad morzem."),
        ("3", "Bocian jest ciężki i musi szybować na prądach ciepłego powietrza, a nad morzem takich prądów prawie nie ma (informacje z 2. akapitu)."),
        ("4", "F, F (młode często startują przed dorosłymi), ?"),
        ("5", "Żeby wyjaśnić trudną rzecz czymś, co czytelnik zna."),
        ("6", "Dwa spośród: zmęczenie, brak pokarmu, zgubienie drogi, przelot nad pustynią lub morzem, samotna droga młodych bocianów."),
        ("7", "zimowy, ptasi, afrykański (mała litera), morski, słoneczny"),
        ("8", "żaby, dżdżownice, przewody, wyruszają, morze, żółty"),
        ("9", "ptaków, ptakom, ptakach"),
        ("10", "Zimą brakuje owadów, więc jaskółki odlatują do Afryki. Jaskółki odlatują do Afryki, ponieważ zimą brakuje owadów. Zwrócić uwagę na odwrotną kolejność zdań."),
        ("11", "Na przykład: igła do szycia, igła sosny lub świerku, igła strzykawki."),
        ("12", "Wyjaśnienie zawiera przyczynę (ciężki ptak, szybowanie, brak prądów nad morzem) i skutek (droga nad lądem). Uczeń pisze własnymi słowami."),
    ],
    7: [
        ("1", "Pies. Dowody, np.: trzyma piłkę w zębach, chodzi na spacery, obwąchuje krzaki, ma miskę, Nowy złapał go za ucho, je skarpetkę."),
        ("2", "Niemowlę, nowe dziecko w rodzinie: jest małe, płacze, pachnie mlekiem, leży na kocu, pierwszy raz się śmieje."),
        ("3", "Był zły albo zazdrosny, bo wszyscy zajmowali się dzieckiem, a spacery były krótsze. Pani odgadła częściowo - nazwała to samotnością."),
        ("4", "F (od trzech miesięcy), F, ?"),
        ("5", "Pierwszy śmiech dziecka sprawił rodzicom ogromną radość - to był dla nich prezent, choć nikt niczego nie dał."),
        ("6", "Na początku pies jest zazdrosny i czuje się odsunięty, na końcu polubił dziecko i chce się nim opiekować jak starszy brat."),
        ("7", "Wcześniej to on był najważniejszy. Wczoraj zjadł skarpetkę."),
        ("8", "biegną - 3. os., lm, czas teraźniejszy; zjadłem - 1. os., lp, czas przeszły; zostanę - 1. os., lp, czas przyszły"),
        ("9", "Pani znieruchomiała z wrażenia."),
        ("10", "piłkę, piłkę, piłką, kołdrą. Ę po pytaniu kogo? co? (biernik), Ą po pytaniu kim? czym? (narzędnik)."),
        ("11", "Zazdrość; np. zazdrosny, zazdrościć."),
        ("12", "Narrator mówi w 1. osobie, nie zdradza się wprost, ale daje wskazówki, kim jest. 5-6 zdań, jedna sytuacja."),
    ],
}


def build_answer_key():
    doc = Document()
    configure_doc(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.4)
    add_footer(doc, "Klucz odpowiedzi")

    p = doc.add_paragraph(style="Title")
    add_text(p, "Klucz odpowiedzi do arkuszy wyrównawczych", size=19, bold=True, color="000000")
    p = doc.add_paragraph()
    set_paragraph(p, after=8)
    add_text(p, "Klucz służy do wspólnego omawiania zadań. W odpowiedziach otwartych należy uznać inne sformułowania zgodne z tekstem i poleceniem.", size=10.5)

    for idx, ws in enumerate(WORKSHEETS):
        if idx > 0:
            add_page_break(doc)
        add_section(doc, f"Arkusz {ws['number']}  {ws['title']}")
        rows = [(n, ans) for n, ans in ANSWERS[ws["number"]]]
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Cm(1.35)
        table.columns[1].width = Cm(15.9)
        hdr = table.rows[0].cells
        for i, label in enumerate(("Zadanie", "Odpowiedź lub kryterium")):
            set_cell_shading(hdr[i], BLUE)
            set_cell_borders(hdr[i])
            set_cell_margins(hdr[i], top=105, bottom=105)
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph(p, after=0)
            add_text(p, label, size=9.5, bold=True, color="FFFFFF")
        set_repeat_table_header(table.rows[0])
        for r_i, (n, ans) in enumerate(rows):
            cells = table.add_row().cells
            cells[0].width = Cm(1.35)
            cells[1].width = Cm(15.9)
            for cell in cells:
                set_cell_borders(cell)
                set_cell_margins(cell, top=90, bottom=90)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if r_i % 2 == 1:
                    set_cell_shading(cell, "F8FAFC")
            p = cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph(p, after=0)
            add_text(p, n, size=9.4, bold=True, color=BLUE)
            p = cells[1].paragraphs[0]
            set_paragraph(p, after=0, line=1.0)
            add_text(p, ans, size=9.4)

        grid_cols = table._tbl.tblGrid.gridCol_lst
        grid_cols[0].set(qn("w:w"), str(Cm(1.7).twips))
        grid_cols[1].set(qn("w:w"), str(Cm(15.4).twips))
        for row in table.rows:
            for idx, width in enumerate((Cm(1.7), Cm(15.4))):
                cell = row.cells[idx]
                cell.width = width
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(width.twips))
                tc_w.set(qn("w:type"), "dxa")

        add_subsection(doc, "Szybkie omówienie po pracy")
        p = doc.add_paragraph()
        set_paragraph(p, after=0)
        add_text(p, "Poproś uczniów, aby wskazali fragment tekstu potwierdzający odpowiedź, poprawili jedno zadanie i nazwali strategię, która im pomogła.", size=9.7)

    path = OUT / "klucz-odpowiedzi-arkusze-wyrownawcze.docx"
    doc.save(path)
    return path


A4_W, A4_H = 841.89, 595.28


def content_clip(page):
    """Prostokat z trescia strony (bez pustych marginesow), zeby A5 wyszlo jak najwieksze."""
    import pymupdf

    rect = pymupdf.Rect()
    for block in page.get_text("blocks"):
        if "język polski" not in block[4]:  # stopka zostaje tylko w A4
            rect |= pymupdf.Rect(block[:4])
    for drawing in page.get_drawings():
        rect |= drawing["rect"]
    return rect + (-6, -6, 6, 6)


def impose_a5(pdf_path, out_path):
    """Dwie strony arkusza jako A5 obok siebie na jednej stronie A4 (poziomo), druk jednostronny."""
    import pymupdf

    src = pymupdf.open(pdf_path)
    if len(src) != 2:
        raise SystemExit(f"{pdf_path.name}: {len(src)} stron, a arkusz ma miec 2")
    out = pymupdf.open()
    page = out.new_page(width=A4_W, height=A4_H)
    margin = 16
    box_w, box_h = A4_W / 2 - 2 * margin, A4_H - 2 * margin
    clips = [content_clip(src[i]) for i in range(2)]
    # Jedna skala dla obu polowek, zeby tekst po lewej i po prawej mial ten sam rozmiar.
    scale = min(min(box_w / c.width, box_h / c.height) for c in clips)
    for i, clip in enumerate(clips):
        x0 = i * A4_W / 2 + margin
        w, h = clip.width * scale, clip.height * scale
        target = pymupdf.Rect(x0 + (box_w - w) / 2, margin, x0 + (box_w + w) / 2, margin + h)
        page.show_pdf_page(target, src, i, clip=clip)
    x = A4_W / 2
    page.draw_line((x, 0), (x, A4_H), color=(0.6, 0.6, 0.6), width=0.4, dashes="[4 3] 0")
    out.save(out_path, garbage=3, deflate=True)


def render_pdfs():
    """Arkusze -> PDF A4 przez headless Chrome (tools/arkusze-pdf.mjs).

    Word na tym komputerze jest bez licencji i zawiesza sie na eksporcie, wiec PDF
    sklada HTML z tych samych danych co DOCX."""
    import json
    import subprocess

    data = [{**ws, "stem": file_stem(ws), "answers": dict(ANSWERS[ws["number"]])} for ws in WORKSHEETS]
    (OUT / "arkusze.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    subprocess.run(["node", str(ROOT / "tools" / "arkusze-pdf.mjs"), str(OUT)], check=True)
    # Ta sama tresc z odpowiedziami do wyswietlenia na projektorze.
    subprocess.run(["node", str(ROOT / "tools" / "arkusze-projektor.mjs"), str(OUT)], check=True)
    return OUT / "pdf"


if __name__ == "__main__":
    import shutil

    made = [build_worksheet(ws) for ws in WORKSHEETS]
    made.append(build_answer_key())
    pdf_dir = render_pdfs()
    public = ROOT / "public" / "materialy" / "wyrownawcze"
    for path in made:
        shutil.copy(path, public / path.name)
    for ws in WORKSHEETS:
        stem = file_stem(ws)
        impose_a5(pdf_dir / f"{stem}.pdf", public / f"{stem}-A5.pdf")
        print(public / f"{stem}-A5.pdf")
